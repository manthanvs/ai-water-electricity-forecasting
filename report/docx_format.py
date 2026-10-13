"""
docx_format.py - formatting primitives for the PCCOE MCA Seminar Report.

Encodes every rule in docs/reference/Seminar_Report_Format_2526.docx:
  * A4, margins Top 25.4 / Bottom 37 / Right 32 / Left 25.4 mm
  * Body Times New Roman 12 pt, justified, 1.5 line spacing,
    first-line indent (five-space convention = 0.5 in), blank line between paragraphs
  * Chapter titles 14 pt bold ALL CAPS centred
  * Section headings 12 pt bold ALL CAPS left; sub-headings Title Case left
  * Figure caption BELOW the figure, table caption ABOVE the table, both 10 pt;
    legends 10 pt below the caption
  * Footer "College Name - Department Name - Academic Year", Times New Roman 10 pt, centred

The cover page and certificate are NOT rebuilt: load_base() opens the official
template, deletes the guideline page and the CONTENTS page, and fills the
placeholders in place, so logos and layout stay identical to the department file.

Tested with python-docx 1.2.0, pdfplumber 0.11, LibreOffice 7+.
Run `python report/docx_format.py --selftest` to regenerate a sample and check it.
"""
from __future__ import annotations

import os
import re
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)

PAGE_W, PAGE_H = Mm(210), Mm(297)
MARGINS = dict(top=Mm(25.4), bottom=Mm(37), right=Mm(32), left=Mm(25.4))
TEXT_WIDTH_MM = 210 - 25.4 - 32  # 152.6 mm usable width

BODY_PT = 12
LINE_SPACING = 1.5
PARA_GAP_PT = 18          # "double space between each paragraph" = one blank 1.5-spaced line
FIRST_INDENT = Inches(0.5)  # "five-space indentation"
CAPTION_PT = 10
FOOTER_PT = 10


# --------------------------------------------------------------------------- low-level helpers
def _ptext(p_el) -> str:
    return "".join(t.text or "" for t in p_el.iter(qn("w:t")))


def _font(run, size=BODY_PT, bold=None, italic=None, color=BLACK, superscript=None, highlight=False):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if superscript is not None:
        run.font.superscript = superscript
    run.font.color.rgb = color
    # strip any highlight inherited from the template
    hl = rpr.find(qn("w:highlight"))
    if hl is not None:
        rpr.remove(hl)
    if highlight:  # used for [[CONFIRM: ...]] markers only
        h = OxmlElement("w:highlight")
        h.set(qn("w:val"), "yellow")
        rpr.append(h)
    return run


def _outline(paragraph, level: int):
    """Outline level gives Word's navigation pane and PDF bookmarks without relying on template styles."""
    ppr = paragraph._p.get_or_add_pPr()
    el = OxmlElement("w:outlineLvl")
    el.set(qn("w:val"), str(level))
    ppr.append(el)


def _spacing(p, before=0, after=0, line=LINE_SPACING):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def _replace_text(p, text, size=None, bold=None):
    """Replace a paragraph's text but keep the first run's formatting; force black, no highlight."""
    runs = p.runs
    if not runs:
        r = p.add_run(text)
        _font(r, size or BODY_PT, bold=bold)
        return
    first = runs[0]
    first.text = text
    for r in runs[1:]:
        r._element.getparent().remove(r._element)
    _font(first, size or (first.font.size.pt if first.font.size else BODY_PT), bold=bold)


def _set_page(section):
    section.page_width, section.page_height = PAGE_W, PAGE_H
    section.top_margin = MARGINS["top"]
    section.bottom_margin = MARGINS["bottom"]
    section.left_margin = MARGINS["left"]
    section.right_margin = MARGINS["right"]
    section.header_distance = Mm(12.5)
    section.footer_distance = Mm(15)


def _page_field(paragraph, size=FOOTER_PT):
    """PAGE field split into separate runs, each carrying Times New Roman, so LibreOffice
    and Word both render the number in the right font (a single-run field falls back to Calibri)."""
    for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), ("result", "1"), ("end", None)):
        run = paragraph.add_run()
        _font(run, size)
        if kind in ("begin", "separate", "end"):
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            run._r.append(fc)
        elif kind is None:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            run._r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)
    # paragraph-mark font as well
    ppr = paragraph._p.get_or_add_pPr()
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(attr), FONT)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(rf)
    rpr.append(sz)
    ppr.append(rpr)


def _clear(part):
    for p in list(part.paragraphs):
        p._p.getparent().remove(p._p)
    for t in list(part.tables):
        t._tbl.getparent().remove(t._tbl)


# --------------------------------------------------------------------------- base document
def load_base(template_path: str | Path, meta: dict) -> Document:
    """
    Open the official template, keep only the cover page + certificate, fill the placeholders.
    meta keys: title, student, prn, guide, program, academic_year, place
    """
    doc = Document(str(template_path))
    body = doc.element.body
    children = list(body.iterchildren())

    def idx(pred):
        for i, el in enumerate(children):
            if el.tag == qn("w:p") and pred(_ptext(el).strip()):
                return i
        raise ValueError("anchor not found in template")

    start = idx(lambda t: t == "Seminar Report")
    end = idx(lambda t: t == "CONTENTS")
    hod = idx(lambda t: t.startswith("Head of the Department"))

    # 1) drop the guideline page (incl. its margin table and section breaks)
    for el in children[:start]:
        body.remove(el)
    # 2) drop everything after the certificate's last line (empty lines + CONTENTS list); keep final sectPr
    for el in children[hod + 1:]:
        if el.tag != qn("w:sectPr"):
            body.remove(el)

    paras = doc.paragraphs
    title_q = "\u201c" + meta["title"] + "\u201d"

    abc_seen = 0
    for p in paras:
        t = p.text.strip()
        if t.startswith("\u201c____") or t.startswith('"____'):
            _replace_text(p, title_q, size=20, bold=True)
        elif t.startswith("Submitted By"):
            _replace_text(p, "Submitted By", size=12, bold=True)
        elif t.replace(".", "").upper() in ("MR ABC",):
            abc_seen += 1
            if abc_seen == 1:
                _replace_text(p, meta["student"], size=12)
            elif abc_seen == 2:
                _replace_text(p, "PRN No. " + meta["prn"], size=12)
            else:
                p._p.getparent().remove(p._p)
        elif t.startswith("Under the Guidance of"):
            _replace_text(p, "Under the Guidance of", size=12, bold=True)
        elif t == "Mr. XYZ":
            _replace_text(p, meta["guide"], size=12)
        elif t.startswith("This is to certify"):
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            segs = [
                ("This is to certify that the seminar report entitled ", False),
                (title_q, True),
                (" being submitted by ", False),
                (f"{meta['student']} (PRN No. {meta['prn']})", True),
                (" is a record of bonafide work carried out by him under the supervision and guidance of ", False),
                (meta["guide"], True),
                (" in partial fulfillment of the requirement for ", False),
                (meta["program"], True),
                (" course of Savitribai Phule Pune University, Pune in the academic year "
                 f"{meta['academic_year']}.", False),
            ]
            for text, b in segs:
                _font(p.add_run(text), 12, bold=b)
            p.paragraph_format.first_line_indent = FIRST_INDENT
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif t.startswith("Place:"):
            _replace_text(p, "Place: " + meta.get("place", "Pune"), size=12)
        elif t.startswith("Name of the Seminar guide"):
            _replace_text(p, meta["guide"] + " (Seminar Guide)", size=12)

    for s in doc.sections:
        _set_page(s)
        _clear(s.footer)
        _clear(s.header)
    return doc


# --------------------------------------------------------------------------- sections / footer
def new_section(doc, meta: dict, number_format: str = "decimal", start: int = 1):
    """
    Start a new page + section with its own footer and page numbering.
    number_format: 'lowerRoman' for front matter, 'decimal' for chapters.
    Footer line 1: College - Department - Academic Year (TNR 10, centred). Line 2: page number.
    """
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    _set_page(sec)
    sec.footer.is_linked_to_previous = False
    sec.header.is_linked_to_previous = False
    _clear(sec.header)
    footer = sec.footer
    _clear(footer)
    p1 = footer.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(p1, 0, 0, 1.0)
    _font(p1.add_run(f"{meta['college']} \u2013 {meta['department']} \u2013 {meta['academic_year']}"), FOOTER_PT)
    p2 = footer.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(p2, 2, 0, 1.0)
    _page_field(p2)
    pg = sec._sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sec._sectPr.append(pg)
    pg.set(qn("w:fmt"), number_format)
    pg.set(qn("w:start"), str(start))
    return sec


def page_break(doc):
    p = doc.add_paragraph()
    _spacing(p, 0, 0, 1.0)
    p.add_run().add_break(WD_BREAK.PAGE)


# --------------------------------------------------------------------------- headings
def front_title(doc, text, new_page: bool = False):
    """ABSTRACT, ACKNOWLEDGEMENT, TABLE OF CONTENTS, ... (same look as a chapter title).
    new_page=True starts it on a new page with 'page break before' (no empty page if the previous one is full)."""
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = new_page
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(p, 0, 24, 1.0)
    _font(p.add_run(text.upper()), 14, bold=True)
    _outline(p, 0)
    return p


def chapter_title(doc, number: int, title: str, new_page: bool = False):
    """'CHAPTER n' and the title; new_page=True puts 'page break before' on it (each chapter starts a new page)."""
    p1 = doc.add_paragraph()
    p1.paragraph_format.page_break_before = new_page
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(p1, 0, 6, 1.0)
    p1.paragraph_format.keep_with_next = True
    _font(p1.add_run(f"CHAPTER {number}"), 14, bold=True)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(p2, 0, 24, 1.0)
    p2.paragraph_format.keep_with_next = True
    _font(p2.add_run(title.upper()), 14, bold=True)
    _outline(p2, 0)


def section_heading(doc, number: str, text: str):
    """'1.1 BACKGROUND' - 12 pt bold ALL CAPS, left; two blank lines' worth of space before."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _spacing(p, 24, 6, 1.0)
    p.paragraph_format.keep_with_next = True
    _font(p.add_run(f"{number} {text.upper()}"), 12, bold=True)
    _outline(p, 1)


def subheading(doc, number: str, text: str):
    """'1.1.1 Data Sources' - 12 pt bold Title Case, left; one blank line before."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _spacing(p, 12, 6, 1.0)
    p.paragraph_format.keep_with_next = True
    _font(p.add_run(f"{number} {text}"), 12, bold=True)
    _outline(p, 2)


# --------------------------------------------------------------------------- body
Segments = list  # list of (text, {"bold":bool,"italic":bool,"superscript":bool,"highlight":bool})


def body_para(doc, content, indent_first=True):
    """content: plain str or list of (text, style-dict) segments."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _spacing(p, 0, PARA_GAP_PT, LINE_SPACING)
    if indent_first:
        p.paragraph_format.first_line_indent = FIRST_INDENT
    segs = [(content, {})] if isinstance(content, str) else content
    for text, st in segs:
        r = _font(p.add_run(text), BODY_PT, bold=st.get("bold", False), italic=st.get("italic", False),
                  superscript=st.get("superscript", None), highlight=st.get("highlight", False))
        if st.get("sub"):
            r.font.subscript = True
    return p


def bullet_list(doc, items, numbered=False):
    """Manual markers + hanging indent (renders identically in Word and LibreOffice)."""
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        last = i == len(items)
        _spacing(p, 0, PARA_GAP_PT if last else 4, LINE_SPACING)
        pf = p.paragraph_format
        pf.left_indent = Inches(0.75)
        pf.first_line_indent = Inches(-0.25)
        pf.tab_stops.add_tab_stop(Inches(0.75))
        marker = f"{i}." if numbered else "\u2022"
        segs = [(item, {})] if isinstance(item, str) else item
        _font(p.add_run(marker + "\t"), BODY_PT)
        for text, st in segs:
            r = _font(p.add_run(text), BODY_PT, bold=st.get("bold", False), italic=st.get("italic", False),
                      superscript=st.get("superscript", None), highlight=st.get("highlight", False))
            if st.get("sub"):
                r.font.subscript = True


def equation(doc, content, label: str):
    """
    Centred equation with right-aligned number, e.g. label='5.1'.
    content: str, or segments [(text, {"sub":True}|{"sup":True}|{"italic":False})] so subscripts are
    real Word subscripts. Do NOT use Unicode subscript letters such as U+209C - Times New Roman lacks them.
    Example: [("MAE = (1/n) Σ |y", {}), ("t", {"sub": True}), (" − ŷ", {}), ("t", {"sub": True}), ("|", {})]
    """
    p = doc.add_paragraph()
    _spacing(p, 6, 12, 1.0)
    tw = Mm(TEXT_WIDTH_MM)
    p.paragraph_format.tab_stops.add_tab_stop(int(tw / 2), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(tw, WD_TAB_ALIGNMENT.RIGHT)
    _font(p.add_run("\t"), BODY_PT)
    segs = [(content, {})] if isinstance(content, str) else content
    for text, st in segs:
        r = _font(p.add_run(text), BODY_PT, italic=st.get("italic", True))
        if st.get("sub"):
            r.font.subscript = True
        if st.get("sup"):
            r.font.superscript = True
    _font(p.add_run(f"\t({label})"), BODY_PT)


def figure(doc, image_path, label: str, caption: str, width_mm: float = 130, legend: str | None = None):
    """Image centred; 'Figure 5.1: caption' BELOW in 10 pt; optional legend below caption in 10 pt."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(p, 6, 4, 1.0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(image_path), width=Mm(min(width_mm, TEXT_WIDTH_MM)))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(c, 0, 4 if legend else PARA_GAP_PT, 1.0)
    c.paragraph_format.keep_with_next = bool(legend)  # caption and legend stay on one page
    _font(c.add_run(f"Figure {label}: "), CAPTION_PT, bold=True)
    _font(c.add_run(caption), CAPTION_PT)
    if legend:
        lg = doc.add_paragraph()
        lg.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _spacing(lg, 0, PARA_GAP_PT, 1.0)
        _font(lg.add_run(legend), CAPTION_PT, italic=True)


def _cell_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), "000000")
        borders.append(e)
    tblPr.append(borders)


def _shade(cell, fill="D9D9D9"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def table(doc, label: str, caption: str, header: list[str], rows: list[list[str]],
          col_widths_mm: list[float], font_pt: float = 10, legend: str | None = None):
    """'Table 5.1: caption' ABOVE in 10 pt; grid table; grey header row (prints cleanly in B&W)."""
    assert abs(sum(col_widths_mm) - TEXT_WIDTH_MM) < 1.0 or sum(col_widths_mm) < TEXT_WIDTH_MM, \
        "table wider than text block"
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(c, 6, 2 if legend else 6, 1.0)
    c.paragraph_format.keep_with_next = True
    _font(c.add_run(f"Table {label}: "), CAPTION_PT, bold=True)
    _font(c.add_run(caption), CAPTION_PT)
    if legend:
        lg = doc.add_paragraph()
        lg.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _spacing(lg, 0, 6, 1.0)
        lg.paragraph_format.keep_with_next = True
        _font(lg.add_run(legend), CAPTION_PT, italic=True)
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, w in enumerate(col_widths_mm):  # grid widths: LibreOffice ignores cell widths without them
        t.columns[i].width = Mm(w)
    _cell_borders(t)
    for r_i, tr in enumerate(t.rows):  # rows never break across pages; header repeats on a new page
        trPr = tr._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        if r_i == 0:
            trPr.append(OxmlElement("w:tblHeader"))
    for r_i, row in enumerate([header] + rows):
        for c_i, val in enumerate(row):
            cell = t.cell(r_i, c_i)
            cell.width = Mm(col_widths_mm[c_i])
            para = cell.paragraphs[0]
            _spacing(para, 1, 1, 1.0)
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            _font(para.add_run(str(val)), font_pt, bold=(r_i == 0))
            if r_i == 0:
                _shade(cell)
    gap = doc.add_paragraph()
    _spacing(gap, 0, PARA_GAP_PT - 6, 1.0)
    return t


def toc_line(doc, text: str, page: str, level: int = 0, bold: bool = False):
    """Right-aligned page number with dot leader; level indents 0.3 in per level."""
    p = doc.add_paragraph()
    _spacing(p, 0, 4, 1.15)
    p.paragraph_format.left_indent = Inches(0.3 * level)
    p.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_WIDTH_MM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    _font(p.add_run(text), BODY_PT, bold=bold)
    _font(p.add_run("\t" + str(page)), BODY_PT, bold=bold)


def reference(doc, number: int, text):
    """IEEE reference with hanging number. text: str, or segments [(text, {"italic": True})] for italic venues."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _spacing(p, 0, 8, 1.15)
    p.paragraph_format.left_indent = Inches(0.45)
    p.paragraph_format.first_line_indent = Inches(-0.45)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(0.45))
    _font(p.add_run(f"[{number}]\t"), BODY_PT)
    for seg, st in ([(text, {})] if isinstance(text, str) else text):
        _font(p.add_run(seg), BODY_PT, italic=st.get("italic", False))


def group_heading(doc, text: str):
    """Unnumbered bold group label inside a front-matter list (e.g. 'Abbreviations', 'Symbols')."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _spacing(p, 12, 4, 1.0)
    p.paragraph_format.keep_with_next = True
    _font(p.add_run(text), BODY_PT, bold=True)
    return p


def term_line(doc, term, meaning):
    """Nomenclature entry: term in bold, meaning after a tab, hanging indent. term/meaning: str or segments."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _spacing(p, 0, 4, 1.15)
    p.paragraph_format.left_indent = Inches(1.4)
    p.paragraph_format.first_line_indent = Inches(-1.4)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(1.4))
    for seg, st in ([(term, {})] if isinstance(term, str) else term):
        r = _font(p.add_run(seg), BODY_PT, bold=True, italic=st.get("italic", False))
        if st.get("sub"):
            r.font.subscript = True
    _font(p.add_run("\t"), BODY_PT)
    for seg, st in ([(meaning, {})] if isinstance(meaning, str) else meaning):
        r = _font(p.add_run(seg), BODY_PT, italic=st.get("italic", False))
        if st.get("sub"):
            r.font.subscript = True


# --------------------------------------------------------------------------- PDF + two-pass page numbers
def soffice_bin() -> str:
    cand = os.environ.get("SOFFICE") or shutil.which("soffice") or shutil.which("libreoffice")
    win = r"C:\Program Files\LibreOffice\program\soffice.exe"
    if not cand and os.path.exists(win):
        cand = win
    if not cand:
        raise RuntimeError("LibreOffice not found - install it or set SOFFICE=/path/to/soffice")
    return cand


def to_pdf(docx_path: str | Path) -> Path:
    docx_path = Path(docx_path)
    subprocess.run([soffice_bin(), "--headless", "--convert-to", "pdf", "--outdir",
                    str(docx_path.parent), str(docx_path)], check=True, capture_output=True)
    return docx_path.with_suffix(".pdf")


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip().lower()


def page_map(pdf_path, entries):
    """
    entries: list of dicts {key, needle, zone} where zone is 'front' or 'body'.
    Returns {key: printed_page_label}. The printed label is read from the last line of each
    page (the footer page number), so roman/arabic numbering is handled automatically.
    'body' entries are searched only after the page that contains 'CHAPTER 1' as a heading line,
    so the TOC / list pages never produce false matches.
    """
    import pdfplumber
    with pdfplumber.open(str(pdf_path)) as pdf:
        texts = [pg.extract_text() or "" for pg in pdf.pages]
    labels = []
    for t in texts:
        lines = [ln.strip() for ln in t.splitlines() if ln.strip()]
        labels.append(lines[-1] if lines and re.fullmatch(r"[ivxlcdm]+|\d+", lines[-1]) else "")
    body_start = 0
    for i, t in enumerate(texts):
        if any(ln.strip() == "CHAPTER 1" for ln in t.splitlines()) and labels[i] == "1":
            body_start = i
            break
    out = {}
    for e in entries:
        lo = body_start if e["zone"] == "body" else 0
        needle = _norm(e["needle"])
        out[e["key"]] = "?"
        for i in range(lo, len(texts)):
            if needle in _norm(texts[i]):
                out[e["key"]] = labels[i] or "?"
                break
    return out


def page_count(pdf_path) -> tuple[int, int]:
    """(total pages, body pages = pages carrying an arabic page number)."""
    import pdfplumber
    with pdfplumber.open(str(pdf_path)) as pdf:
        total = len(pdf.pages)
        body = 0
        for pg in pdf.pages:
            lines = [ln.strip() for ln in (pg.extract_text() or "").splitlines() if ln.strip()]
            if lines and lines[-1].isdigit():
                body += 1
    return total, body


# --------------------------------------------------------------------------- self-test
def _selftest(out_dir: Path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np

    here = Path(__file__).resolve().parent
    template = here.parent / "docs" / "reference" / "Seminar_Report_Format_2526.docx"
    out_dir.mkdir(parents=True, exist_ok=True)

    meta = dict(
        title="AI-Based Prediction of Water Demand and Electricity Consumption",
        student="Mr. Manthan Sankpal", prn="125M1H064",
        guide="Prof. Dr. Anjana Arakerimath",
        program="Master of Computer Applications (MCA)",
        academic_year="2026-27", place="Pune",
        college="Pimpri Chinchwad College of Engineering",
        department="Department of Master of Computer Application",
    )

    # grayscale sample figure (style the real figures the same way)
    x = np.arange(168)
    y = 2 + np.sin(2 * np.pi * x / 24) + 0.3 * np.sin(2 * np.pi * x / 168)
    fig, ax = plt.subplots(figsize=(6.0, 2.6), dpi=300)
    ax.plot(x, y, color="black", lw=1.0, ls="-", label="Actual")
    ax.plot(x, y * 0.97 + 0.05, color="0.45", lw=1.0, ls="--", label="Forecast")
    ax.set_xlabel("Hour"); ax.set_ylabel("kW"); ax.legend(frameon=False); ax.grid(alpha=0.3, color="0.7")
    fig.tight_layout(); img = out_dir / "selftest_fig.png"; fig.savefig(img); plt.close(fig)

    def build(pages):
        d = load_base(template, meta)
        new_section(d, meta, "lowerRoman", 1)
        front_title(d, "Abstract")
        body_para(d, "Sample abstract paragraph. " * 12)
        page_break(d)
        front_title(d, "Table of Contents")
        toc_line(d, "Abstract", pages.get("abs", "?"))
        toc_line(d, "Chapter 1  Introduction", pages.get("c1", "?"), bold=True)
        toc_line(d, "1.1 Background", pages.get("s11", "?"), level=1)
        toc_line(d, "References", pages.get("ref", "?"), bold=True)
        new_section(d, meta, "decimal", 1)
        chapter_title(d, 1, "Introduction")
        section_heading(d, "1.1", "Background")
        body_para(d, "Body text paragraph in Times New Roman 12 pt, justified, 1.5 spacing. " * 6)
        body_para(d, [("A CONFIRM marker renders like this: ", {}),
                      ("[[CONFIRM: Director's name]]", {"highlight": True}), (".", {})])
        subheading(d, "1.1.1", "Smart Meter Data")
        bullet_list(d, ["First point.", "Second point that is long enough to wrap onto a second line "
                        "so the hanging indent can be checked visually."])
        equation(d, [("MAE = (1/n) \u03a3 |y", {}), ("t", {"sub": True}), (" \u2212 \u0177", {}),
                     ("t", {"sub": True}), ("|", {})], "1.1")
        figure(d, img, "1.1", "Actual versus forecast hourly load (sample)", legend="Solid: actual; dashed: forecast")
        table(d, "1.1", "Sample metrics table", ["Model", "MAE", "RMSE", "MAPE (%)"],
              [["Seasonal naive", "0.41", "0.55", "18.2"], ["LSTM", "0.29", "0.40", "12.7"]],
              [52.6, 33, 33, 34])
        page_break(d)
        front_title(d, "References")
        reference(d, 1, "S. Hochreiter and J. Schmidhuber, \u201cLong short-term memory,\u201d "
                        "Neural Computation, vol. 9, no. 8, pp. 1735\u20131780, 1997.")
        p = out_dir / "format_selftest.docx"
        d.save(str(p))
        return p

    p1 = build({})
    pdf = to_pdf(p1)
    pages = page_map(pdf, [
        dict(key="abs", needle="ABSTRACT", zone="front"),
        dict(key="c1", needle="CHAPTER 1", zone="body"),
        dict(key="s11", needle="1.1 BACKGROUND", zone="body"),
        dict(key="ref", needle="REFERENCES", zone="body"),
    ])
    p2 = build(pages)
    pdf = to_pdf(p2)
    total, body = page_count(pdf)
    print("page map:", pages, "| total pages:", total, "| body pages:", body)
    assert "?" not in pages.values(), "two-pass page lookup failed"
    print("SELFTEST OK ->", p2, pdf)


if __name__ == "__main__":
    import sys
    if "--selftest" in sys.argv:
        _selftest(Path(__file__).resolve().parent / "out")
